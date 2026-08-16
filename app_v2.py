import os
import shutil
import threading

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")  # no GUI backend on a server
import matplotlib.pyplot as plt
import seaborn as sns

import fitz
from docx import Document
from flask import Flask, request, send_file, jsonify, render_template
from flask_cors import CORS
from fpdf import FPDF
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity
from werkzeug.utils import secure_filename

app = Flask(__name__)
CORS(app)

UPLOAD_FOLDER = "uploads"
COURSE_FOLDER = os.path.join(UPLOAD_FOLDER, "allfiles")
REPORT_DIR = os.path.join("static", "reports")
OUTPUT_CSV = "all_course_similarity_matrix.csv"
RECOMMENDATION_PDF = "recommendation_report.pdf"

# Tuning knobs — named instead of buried as magic numbers
CHUNK_WORDS = 200          # must stay under the model's 256 word-piece limit
CHUNK_OVERLAP = 50
LOW_SCORE_THRESHOLD = 0.15
MAX_HEATMAP_COURSES = 15
RESERVED_COLUMNS = {"SAIL PNO", "Best Match Course", "Highest Similarity Score"}

os.makedirs(COURSE_FOLDER, exist_ok=True)
os.makedirs(REPORT_DIR, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

model = SentenceTransformer("all-MiniLM-L6-v2")

# Output paths are global, so serialise runs to stop two users corrupting
# each other's files. For real multi-user use, key everything by a session id.
run_lock = threading.Lock()


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------

def read_tabular(path, **kwargs):
    """Read .csv or .xlsx transparently. FIX: /upload accepts both, but the
    original code always called read_excel, which crashes on a CSV."""
    if path.lower().endswith(".csv"):
        return pd.read_csv(path, **kwargs)
    return pd.read_excel(path, **kwargs)


def extract_text(file_path):
    try:
        if file_path.lower().endswith(".pdf"):
            with fitz.open(file_path) as doc:
                return "\n".join(page.get_text() for page in doc)
        elif file_path.lower().endswith(".docx"):
            doc = Document(file_path)
            return "\n".join(para.text for para in doc.paragraphs)
    except Exception as e:
        print(f"Error reading {file_path}: {e}")
    return ""


def chunk_text(text, size=CHUNK_WORDS, overlap=CHUNK_OVERLAP):
    """FIX: all-MiniLM-L6-v2 truncates at 256 word-pieces. Embedding a whole
    course document meant everything after roughly page one was discarded
    silently. Chunking preserves the full document."""
    words = text.split()
    if not words:
        return []
    step = max(1, size - overlap)
    chunks = []
    for i in range(0, len(words), step):
        window = words[i:i + size]
        if window:
            chunks.append(" ".join(window))
        if i + size >= len(words):
            break
    return chunks


def latin1_safe(text):
    """FPDF1 only speaks Latin-1; a curly quote or en-dash in a filename
    otherwise raises UnicodeEncodeError mid-report."""
    return str(text).encode("latin-1", "replace").decode("latin-1")


def unique_key(name, taken):
    """Stop same-named files (or a file colliding with a reserved column)
    from silently overwriting each other."""
    candidate = name
    n = 2
    while candidate in taken or candidate in RESERVED_COLUMNS:
        candidate = f"{name} ({n})"
        n += 1
    return candidate


# --------------------------------------------------------------------------
# Routes
# --------------------------------------------------------------------------

@app.route("/")
def login_page():
    return render_template("index.html")


@app.route("/register")
def register():
    return render_template("register.html")


@app.route("/dashboard")
def upload_page():
    return render_template("dashboard.html")


@app.route("/upload", methods=["POST"])
def upload_files():
    print("Received upload request")

    csv_file = request.files.get("csv")
    if not csv_file or not csv_file.filename:
        return jsonify({"success": False, "error": "Missing CSV/XLSX file."}), 400

    ext = os.path.splitext(csv_file.filename)[1].lower()
    if ext not in [".csv", ".xlsx"]:
        return jsonify({"success": False,
                        "error": "Unsupported file format. Upload CSV or XLSX only."}), 400

    # FIX: MultiDict.items() returns only the FIRST value per key. If the
    # frontend appends every file under the same key, all but one were
    # silently dropped. multi=True yields all of them.
    folder_files = [f for k, f in request.files.items(multi=True)
                    if k.startswith("folder")]
    folder_files = [f for f in folder_files
                    if f.filename and f.filename.lower().endswith((".pdf", ".docx"))]

    if not folder_files:
        return jsonify({"success": False, "error": "No course documents uploaded."}), 400

    with run_lock:
        shutil.rmtree(UPLOAD_FOLDER, ignore_errors=True)
        os.makedirs(COURSE_FOLDER, exist_ok=True)

        csv_path = os.path.join(UPLOAD_FOLDER, "requirement" + ext)
        csv_file.save(csv_path)
        print(f"Saved training need file: {csv_path}")

        for f in folder_files:
            f.save(os.path.join(COURSE_FOLDER, secure_filename(f.filename)))
        print(f"Saved {len(folder_files)} course documents")

        try:
            print("Running similarity computation...")
            run_similarity(csv_path, COURSE_FOLDER, OUTPUT_CSV)
        except Exception as e:
            print(f"Error during similarity run: {e}")
            return jsonify({"success": False, "error": str(e)}), 500

    return jsonify({"success": True})


@app.route("/success")
def success_page():
    return render_template("success.html")


@app.route("/download")
def download_file():
    if not os.path.exists(OUTPUT_CSV):
        return jsonify({"success": False,
                        "error": "No results yet. Run an analysis first."}), 404
    return send_file(OUTPUT_CSV, as_attachment=True)


@app.route("/download-report")
def download_report():
    if not os.path.exists(RECOMMENDATION_PDF):
        return jsonify({"success": False,
                        "error": "No report yet. Run an analysis first."}), 404
    return send_file(RECOMMENDATION_PDF, as_attachment=True)


# --------------------------------------------------------------------------
# Core logic
# --------------------------------------------------------------------------

def run_similarity(file_path, course_folder, output_csv):
    df_all = read_tabular(file_path, header=None)

    header_row = None
    for i in range(len(df_all)):
        row = df_all.iloc[i].astype(str)
        if row.str.contains("SAIL", case=False, na=False).any() and \
           row.str.contains("DEVELOPMENT", case=False, na=False).any():
            header_row = i
            break

    if header_row is None:
        raise ValueError("Could not locate the header row with 'SAIL PNO' and "
                         "'DEVELOPMENT NEEDS'.")

    df = read_tabular(file_path, header=header_row)
    # collapse internal whitespace too, so 'SAIL  PNO' still matches
    df = df.rename(columns=lambda x: " ".join(str(x).split()).upper())

    if "SAIL PNO" not in df.columns or "DEVELOPMENT NEEDS" not in df.columns:
        raise ValueError("Required columns 'SAIL PNO' and 'DEVELOPMENT NEEDS' "
                         f"not found. Found: {list(df.columns)}")

    df = df[["SAIL PNO", "DEVELOPMENT NEEDS"]].dropna()
    if df.empty:
        raise ValueError("No usable rows found after dropping blanks.")

    pnos = df["SAIL PNO"].astype(str).tolist()
    needs = df["DEVELOPMENT NEEDS"].astype(str).tolist()

    # ---- Load course documents -------------------------------------------
    course_texts = {}
    for root, _, files in os.walk(course_folder):
        for file in sorted(files):
            if not file.lower().endswith((".pdf", ".docx")):
                continue
            text = extract_text(os.path.join(root, file))
            if not text.strip():
                print(f"Skipped (no extractable text): {file}")
                continue
            base = os.path.basename(file).replace("All_designs_", "")
            course_texts[unique_key(base, course_texts)] = text

    if not course_texts:
        raise ValueError("No valid course documents found.")

    course_names = list(course_texts.keys())

    # ---- Chunk, then embed -----------------------------------------------
    all_chunks, chunk_owner = [], []
    for name, text in course_texts.items():
        pieces = chunk_text(text)
        if not pieces:
            continue
        all_chunks.extend(pieces)
        chunk_owner.extend([name] * len(pieces))

    if not all_chunks:
        raise ValueError("Course documents produced no text chunks.")

    print(f"Embedding {len(needs)} needs and {len(all_chunks)} course chunks "
          f"from {len(course_names)} documents...")

    need_emb = model.encode(needs, batch_size=64, show_progress_bar=False)
    chunk_emb = model.encode(all_chunks, batch_size=64, show_progress_bar=False)
    need_emb = np.asarray(need_emb)
    chunk_emb = np.asarray(chunk_emb)

    # One matrix op instead of a per-row Python loop
    sims = cosine_similarity(need_emb, chunk_emb)   # (n_needs, n_chunks)

    owner = np.asarray(chunk_owner)
    per_course = np.zeros((len(needs), len(course_names)), dtype=float)
    for j, name in enumerate(course_names):
        mask = owner == name
        if mask.any():
            # best-matching passage represents the document
            per_course[:, j] = sims[:, mask].max(axis=1)

    # ---- Assemble output --------------------------------------------------
    best_idx = per_course.argmax(axis=1)
    records = []
    for i in range(len(needs)):
        row = {
            "SAIL PNO": pnos[i],
            "Best Match Course": course_names[best_idx[i]],
            "Highest Similarity Score": round(float(per_course[i, best_idx[i]]), 4),
        }
        for j, name in enumerate(course_names):
            row[name] = round(float(per_course[i, j]), 4)
        records.append(row)

    pd.DataFrame(records).to_csv(output_csv, index=False)
    generate_detailed_pdf_report(records)
    print("Similarity matrix and report saved.")


# --------------------------------------------------------------------------
# Reporting
# --------------------------------------------------------------------------

def _bar_chart(labels, scores, color, title, path):
    plt.figure(figsize=(9, 4.5))
    short = [str(l)[:38] + ("..." if len(str(l)) > 38 else "") for l in labels]
    bars = plt.bar(short, scores, color=color)
    plt.xticks(rotation=45, ha="right", fontsize=8)
    plt.ylim(0, 1.0)
    plt.ylabel("Similarity Score")
    plt.title(title)
    for bar in bars:
        h = bar.get_height()
        plt.text(bar.get_x() + bar.get_width() / 2.0, h + 0.01,
                 f"{h:.2f}", ha="center", fontsize=8)
    plt.tight_layout()
    plt.savefig(path, dpi=120)
    plt.close()


def generate_detailed_pdf_report(records, output_pdf=RECOMMENDATION_PDF):
    df = pd.DataFrame(records)
    df["Highest Similarity Score"] = df["Highest Similarity Score"].astype(float).round(4)

    top_5 = df.sort_values("Highest Similarity Score", ascending=False).head(5)
    bottom_5 = df.sort_values("Highest Similarity Score").head(5)
    course_frequency = df["Best Match Course"].value_counts().head(10)
    low_performers = (df[df["Highest Similarity Score"] < LOW_SCORE_THRESHOLD]
                      ["Best Match Course"].value_counts().head(10))

    avg_score = df["Highest Similarity Score"].mean()
    min_score = df["Highest Similarity Score"].min()
    max_score = df["Highest Similarity Score"].max()
    std_score = df["Highest Similarity Score"].std()

    os.makedirs(REPORT_DIR, exist_ok=True)
    images = []

    _bar_chart(top_5["Best Match Course"], top_5["Highest Similarity Score"],
               "seagreen", "Top 5 Courses by Similarity Score",
               os.path.join(REPORT_DIR, "top5.png"))
    images.append(("Top 5 Most Relevant Courses", "top5.png"))

    _bar_chart(bottom_5["Best Match Course"], bottom_5["Highest Similarity Score"],
               "tomato", "Bottom 5 Courses by Similarity Score",
               os.path.join(REPORT_DIR, "bottom5.png"))
    images.append(("Bottom 5 Least Relevant Courses", "bottom5.png"))

    plt.figure(figsize=(9, 4.5))
    short = [str(l)[:38] for l in course_frequency.index]
    plt.bar(short, course_frequency.values, color="royalblue")
    plt.xticks(rotation=45, ha="right", fontsize=8)
    plt.ylabel("Times recommended")
    plt.title("Most Frequently Recommended Courses")
    plt.tight_layout()
    plt.savefig(os.path.join(REPORT_DIR, "frequency.png"), dpi=120)
    plt.close()
    images.append(("Most Frequently Recommended Courses", "frequency.png"))

    plt.figure(figsize=(9, 4.5))
    sns.histplot(df["Highest Similarity Score"], bins=10, kde=True, color="purple")
    plt.title("Similarity Score Distribution")
    plt.tight_layout()
    plt.savefig(os.path.join(REPORT_DIR, "distribution.png"), dpi=120)
    plt.close()
    images.append(("Similarity Score Distribution", "distribution.png"))

    # Heatmap only when it will actually be readable
    numeric = df.drop(columns=["SAIL PNO", "Best Match Course",
                               "Highest Similarity Score"], errors="ignore")
    numeric = numeric.select_dtypes(include="number")
    if 2 <= numeric.shape[1] <= MAX_HEATMAP_COURSES and len(df) > 1:
        plt.figure(figsize=(10, 8))
        sns.heatmap(numeric.corr(), annot=True, fmt=".2f", cmap="coolwarm",
                    annot_kws={"size": 7})
        plt.title("Course Score Correlation")
        plt.tight_layout()
        plt.savefig(os.path.join(REPORT_DIR, "heatmap.png"), dpi=120)
        plt.close()
        images.append(("Score Correlation Heatmap", "heatmap.png"))
    else:
        print(f"Heatmap skipped ({numeric.shape[1]} courses — "
              f"unreadable above {MAX_HEATMAP_COURSES}).")

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Training Recommendation Report", ln=True, align="C")
    pdf.set_font("Arial", "", 12)
    pdf.ln(6)
    pdf.multi_cell(0, 8, latin1_safe(
        "This report summarizes the matching between employee training needs "
        "and available course documents.\n\n"
        f"Total Entries: {len(df)}\n"
        f"Average Score: {avg_score:.4f}\n"
        f"Highest Score: {max_score:.4f}\n"
        f"Lowest Score: {min_score:.4f}\n"
        f"Standard Deviation: {std_score:.4f}"
    ))

    for title, img in images:
        path = os.path.join(REPORT_DIR, img)
        if not os.path.exists(path):
            continue
        pdf.add_page()
        pdf.set_font("Arial", "B", 14)
        pdf.cell(0, 10, latin1_safe(title), ln=True)
        pdf.image(path, x=10, w=190)

    if not low_performers.empty:
        pdf.add_page()
        pdf.set_font("Arial", "B", 14)
        pdf.cell(0, 10, "Courses Recommended for Review", ln=True)
        pdf.set_font("Arial", "", 11)
        for name, count in low_performers.items():
            pdf.multi_cell(0, 8, latin1_safe(
                f"{name}: {count} entries below {LOW_SCORE_THRESHOLD} score"))

    pdf.output(output_pdf)


# --------------------------------------------------------------------------

if __name__ == "__main__":
    print("Starting Flask server...")
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1")