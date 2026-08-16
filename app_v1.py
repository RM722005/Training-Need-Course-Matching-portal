import os
import shutil
from flask import Flask, request, send_file, jsonify, render_template
from flask_cors import CORS
from werkzeug.utils import secure_filename
import pandas as pd
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity
from docx import Document
import fitz
import matplotlib.pyplot as plt
import seaborn as sns
from fpdf import FPDF

app = Flask(__name__)
CORS(app)

UPLOAD_FOLDER = 'uploads'
COURSE_FOLDER = os.path.join(UPLOAD_FOLDER, 'allfiles')
OUTPUT_CSV = 'all_course_similarity_matrix.csv'
RECOMMENDATION_PDF = 'recommendation_report.pdf'

os.makedirs(COURSE_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

model = SentenceTransformer('all-MiniLM-L6-v2')


def extract_text(file_path):
    try:
        if file_path.endswith(".pdf"):
            with fitz.open(file_path) as doc:
                return "\n".join(page.get_text() for page in doc)
        elif file_path.endswith(".docx"):
            doc = Document(file_path)
            return "\n".join(para.text for para in doc.paragraphs)
    except Exception as e:
        print(f"Error reading {file_path}: {e}")
    return ""


@app.route('/')
def login_page():
    return render_template('index.html')


@app.route('/register')
def register():
    return render_template('register.html')


@app.route('/dashboard')
def upload_page():
    return render_template('dashboard.html')


@app.route('/upload', methods=['POST'])
def upload_files():
    print("📥 Received upload request")
    shutil.rmtree(UPLOAD_FOLDER, ignore_errors=True)
    os.makedirs(COURSE_FOLDER, exist_ok=True)

    csv_file = request.files.get('csv')
    if not csv_file:
        return jsonify({"success": False, "error": "Missing CSV/XLSX file."})

    ext = os.path.splitext(csv_file.filename)[1].lower()
    if ext not in ['.csv', '.xlsx']:
        return jsonify({"success": False, "error": "Unsupported file format. Upload CSV or XLSX only."})

    csv_path = os.path.join(UPLOAD_FOLDER, 'requirement' + ext)
    csv_file.save(csv_path)
    print(f"✅ Saved training need file: {csv_path}")

    folder_files = [f for k, f in request.files.items() if k.startswith('folder')]
    if not folder_files:
        return jsonify({"success": False, "error": "No course documents uploaded."})

    for f in folder_files:
        if f.filename.endswith(('.pdf', '.docx')):
            f.save(os.path.join(COURSE_FOLDER, secure_filename(f.filename)))

    try:
        print("🔍 Running similarity computation...")
        run_similarity(csv_path, COURSE_FOLDER, OUTPUT_CSV)
    except Exception as e:
        print(f"❌ Error during similarity run: {e}")
        return jsonify({"success": False, "error": str(e)})

    return jsonify({"success": True})


@app.route('/success')
def success_page():
    return render_template('success.html')


@app.route('/download')
def download_file():
    return send_file(OUTPUT_CSV, as_attachment=True)


@app.route('/download-report')
def download_report():
    return send_file(RECOMMENDATION_PDF, as_attachment=True)


def run_similarity(file_path, course_folder, output_csv):
    df_all = pd.read_excel(file_path, header=None)
    header_row = None

    for i in range(len(df_all)):
        row = df_all.iloc[i]
        if row.astype(str).str.contains("SAIL", case=False).any() and \
           row.astype(str).str.contains("DEVELOPMENT", case=False).any():
            header_row = i
            break

    if header_row is None:
        raise ValueError("Could not locate the header row with 'SAIL PNO' and 'DEVELOPMENT NEEDS'.")

    df = pd.read_excel(file_path, header=header_row)
    df = df.rename(columns=lambda x: str(x).strip().upper())

    if 'SAIL PNO' not in df.columns or 'DEVELOPMENT NEEDS' not in df.columns:
        raise ValueError("Required columns 'SAIL PNO' and 'DEVELOPMENT NEEDS' not found.")

    df = df[['SAIL PNO', 'DEVELOPMENT NEEDS']].dropna()

    pnos = df['SAIL PNO'].astype(str).tolist()
    needs = df['DEVELOPMENT NEEDS'].astype(str).tolist()
    need_embeddings = model.encode(needs, convert_to_tensor=True)

    course_texts = {}
    for root, _, files in os.walk(course_folder):
        for file in files:
            if file.endswith(('.pdf', '.docx')):
                path = os.path.join(root, file)
                text = extract_text(path)
                if text:
                    cleaned_name = os.path.basename(file).replace("All_designs_", "")
                    course_texts[cleaned_name] = text

    if not course_texts:
        raise ValueError("No valid course documents found.")

    course_names = list(course_texts.keys())
    course_embeddings = model.encode(list(course_texts.values()), convert_to_tensor=True)

    similarity_matrix = []
    for i, emb in enumerate(need_embeddings):
        scores = cosine_similarity([emb.cpu().numpy()], course_embeddings.cpu().numpy())[0]
        best_idx = scores.argmax()
        row = {
            "SAIL PNO": pnos[i],
            "Best Match Course": course_names[best_idx],
            "Highest Similarity Score": round(scores[best_idx], 4),
        }
        for j, name in enumerate(course_names):
            row[name] = round(scores[j], 4)
        similarity_matrix.append(row)

    pd.DataFrame(similarity_matrix).to_csv(output_csv, index=False)
    generate_detailed_pdf_report(similarity_matrix)
    print(f"✅ Similarity matrix and report saved.")


def generate_detailed_pdf_report(similarity_matrix, output_pdf="recommendation_report.pdf"):
    df = pd.DataFrame(similarity_matrix)
    df["Highest Similarity Score"] = df["Highest Similarity Score"].astype(float).round(4)

    top_5 = df.sort_values("Highest Similarity Score", ascending=False).head(5)
    bottom_5 = df.sort_values("Highest Similarity Score").head(5)
    course_frequency = df["Best Match Course"].value_counts().head(10)
    low_performers = df[df["Highest Similarity Score"] < 0.15]["Best Match Course"].value_counts().head(10)

    avg_score = df["Highest Similarity Score"].mean()
    min_score = df["Highest Similarity Score"].min()
    max_score = df["Highest Similarity Score"].max()
    std_score = df["Highest Similarity Score"].std()

    os.makedirs("static/reports", exist_ok=True)

    # --- Top 5 Chart ---
    plt.figure(figsize=(8, 4))
    scores = top_5["Highest Similarity Score"]
    labels = top_5["Best Match Course"]
    bars = plt.bar(labels, scores, color='seagreen')
    plt.xticks(rotation=45)
    plt.ylim(0, 1.0)
    plt.ylabel("Similarity Score")
    plt.title("Top 5 Courses by Similarity Score")
    for bar in bars:
        height = bar.get_height()
        plt.text(bar.get_x() + bar.get_width()/2.0, height + 0.01, f"{height:.2f}", ha='center', fontsize=9)
    plt.tight_layout()
    plt.savefig("static/reports/top5.png")
    plt.close()

    # --- Bottom 5 Chart ---
    plt.figure(figsize=(8, 4))
    scores = bottom_5["Highest Similarity Score"]
    labels = bottom_5["Best Match Course"]
    bars = plt.bar(labels, scores, color='tomato')
    plt.xticks(rotation=45)
    plt.ylim(0, 1.0)
    plt.ylabel("Similarity Score")
    plt.title("Bottom 5 Courses by Similarity Score")
    for bar in bars:
        height = bar.get_height()
        plt.text(bar.get_x() + bar.get_width()/2.0, height + 0.01, f"{height:.2f}", ha='center', fontsize=9)
    plt.tight_layout()
    plt.savefig("static/reports/bottom5.png")
    plt.close()

    # --- Course Frequency Chart ---
    course_frequency.plot(kind='bar', color='royalblue', title='Top Recommended Courses')
    plt.xticks(rotation=45)
    plt.tight_layout()
    plt.savefig("static/reports/frequency.png")
    plt.close()

    # --- Score Distribution ---
    sns.histplot(df["Highest Similarity Score"], bins=10, kde=True, color="purple")
    plt.title("Score Distribution")
    plt.tight_layout()
    plt.savefig("static/reports/distribution.png")
    plt.close()

    # --- Correlation Heatmap ---
    plt.figure(figsize=(10, 8))
    sns.heatmap(df.drop(columns=["SAIL PNO"]).corr(numeric_only=True), annot=True, cmap="coolwarm")
    plt.title("Correlation Heatmap")
    plt.tight_layout()
    plt.savefig("static/reports/heatmap.png")
    plt.close()

    # --- PDF Report ---
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, "Training Recommendation Report", ln=True, align='C')
    pdf.set_font("Arial", '', 12)
    pdf.ln(10)
    pdf.multi_cell(0, 10, f"""
This report summarizes the matching between employee training needs and available course documents.

Total Entries: {len(df)}
Average Score: {avg_score:.4f}
Highest Score: {max_score:.4f}
Lowest Score: {min_score:.4f}
Standard Deviation: {std_score:.4f}
""")

    for title, img in [
        ("Top 5 Most Relevant Courses", "top5.png"),
        ("Bottom 5 Least Relevant Courses", "bottom5.png"),
        ("Most Frequently Recommended Courses", "frequency.png"),
        ("Similarity Score Distribution", "distribution.png"),
        ("Score Correlation Heatmap", "heatmap.png")
    ]:
        pdf.add_page()
        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 10, title, ln=True)
        pdf.image(f"static/reports/{img}", x=10, w=190)

    if not low_performers.empty:
        pdf.add_page()
        pdf.set_font("Arial", 'B', 14)
        pdf.cell(0, 10, "Courses Recommended for Review", ln=True)
        pdf.set_font("Arial", '', 11)
        for name, count in low_performers.items():
            pdf.cell(0, 10, f"{name}: {count} entries < 0.15 score", ln=True)

    pdf.output(output_pdf)


# -------------------- START SERVER --------------------

if __name__ == '__main__':
    print("✅ Starting Flask server...")
    app.run(debug=True)